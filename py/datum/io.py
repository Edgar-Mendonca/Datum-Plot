"""Reading data files for File › Import: CSV, TSV, TXT and Excel workbooks. Every cell comes back
as text (numbers without trailing zeros, dates as ISO), the Rust engine parses the numbers."""
import csv
import datetime
import io
import itertools
import os
import re

from .files import atomic

MAX_ROWS, MAX_COLS = 1_000_000, 500
MAX_BYTES = 300 << 20  # a file read into memory at once
DECIMAL_COMMA = re.compile(r"^-?\d+,\d+$")


def cell(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v).upper()
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() and abs(v) < 1e15 else repr(v)
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.isoformat()
    return str(v).strip()


def read_text(path):
    """A delimited text file: the delimiter is guessed; 1,5 becomes 1.5 when ; separates columns."""
    with open(path, "rb") as f:
        raw = f.read()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = "\n".join(text.splitlines()[:50])
    try:
        delim = csv.Sniffer().sniff(sample, delimiters=",;\t| ").delimiter
    except csv.Error:
        delim = "\t" if "\t" in sample else ","
    rows = []
    try:
        lines = list(itertools.islice(csv.reader(io.StringIO(text.replace("\x00", ""), newline=""), delimiter=delim, skipinitialspace=delim == " "), MAX_ROWS))
    except csv.Error as e:
        raise ValueError(f"That file could not be read as text columns ({e}).")
    for r in lines:
        if len(rows) >= MAX_ROWS:
            break
        r = [c.strip() for c in r[:MAX_COLS]]
        if delim == ";":
            r = [c.replace(",", ".") if DECIMAL_COMMA.match(c) else c for c in r]
        rows.append(r)
    while rows and not any(rows[-1]):
        rows.pop()
    return [{"name": os.path.splitext(os.path.basename(path))[0], "rows": rows, "delimiter": delim}]


def read_excel(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(values_only=True, max_col=MAX_COLS):
            if len(rows) >= MAX_ROWS:
                break
            rows.append([cell(v) for v in r])
        while rows and not any(rows[-1]):
            rows.pop()
        width = max((max((i + 1 for i, c in enumerate(r) if c), default=0) for r in rows), default=0)
        out.append({"name": ws.title, "rows": [r[:width] for r in rows]})
    wb.close()
    return out


def read_file(path):
    """{"sheets": [{"name", "rows": [[text]]}]} for a CSV, TSV, TXT or XLSX file."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".pzfx", ".xlsx", ".xlsm", ".csv", ".tsv", ".txt", ".dat", ".prn") and os.path.getsize(path) > MAX_BYTES:
        raise ValueError(f"{os.path.basename(path)} is larger than {MAX_BYTES >> 20} MB; split it into smaller files.")
    if ext == ".pzfx":
        from .imports import read_pzfx
        return {"sheets": [], "tables": read_pzfx(path)}
    if ext in (".xlsx", ".xlsm"):
        try:
            sheets = read_excel(path)
        except (OSError, ValueError, KeyError, EOFError) as e:  # BadZipFile is an OSError's cousin: catch it by name below
            raise ValueError(f"That .xlsx file is damaged or not a workbook ({type(e).__name__}).")
        except Exception as e:
            if type(e).__name__ in ("BadZipFile", "InvalidFileException", "LargeZipFile"):
                raise ValueError(f"That .xlsx file is damaged or not a workbook ({type(e).__name__}).")
            raise
    elif ext in (".csv", ".tsv", ".txt", ".dat", ".prn"):
        sheets = read_text(path)
    else:
        raise ValueError(f"Datum Plot cannot read {ext or 'this kind of'} files yet. Use CSV, TXT, .xlsx or .pzfx.")
    from .imports import plates
    return {"sheets": sheets, "plates": plates(sheets)}


def sheets(project, analyses):
    """The workbook's sheets, whatever the file type: [(name, rows, wide first column)], each cell a value or
    (value, style) with style head, bold, title or excl (an excluded value)."""
    from . import analysis, stats
    out, used = [], set()

    def name_of(name):
        base = re.sub(r"[\[\]:*?/\\]", " ", name).strip()[:28] or "Sheet"
        n, k = base, 2
        while n.lower() in used:
            n, k = f"{base[:25]} ({k})", k + 1
        used.add(n.lower())
        return n

    for t in project.get("tables", []):
        xy = t.get("kind", "xy") in ("xy", "survival", "matrix")
        grouped = t.get("kind") in ("grouped", "contingency", "parts", "multi", "sheet")
        heads = ([f"{t.get('xTitle') or 'X'}{' (' + t['xUnits'] + ')' if t.get('xUnits') else ''}"] if xy else ["Row"] if grouped else [])
        sub = {"msn": ["Mean", "SD", "N"], "mesn": ["Mean", "SEM", "N"]}
        for s, (c0, w) in zip(t["sets"], stats.set_columns(t)):
            heads += [f"{s['title']}" + (f" {sub[s['entry']][j]}" if s["entry"] != "rep" else f" Y{j + 1}" if w > 1 else "") for j in range(w)]
        rows = [[(h, "head") for h in heads]]
        ex = {tuple(e) for e in t.get("excluded", [])}
        first = 0 if xy else 1
        for r in range(t["rows"]):
            row = [t["cols"][c][r] if c < len(t["cols"]) and r < len(t["cols"][c]) else None for c in range(first, len(t["cols"]))]
            if any(v is not None for v in row):
                lead = [(t.get("rowTitles") or [""] * t["rows"])[r] or f"Row {r + 1}"] if grouped else []
                rows.append(lead + [(v, "excl") if (r, j + first) in ex else v for j, v in enumerate(row)])
        out.append((name_of(t["name"]), rows, False))
    for a in analyses:
        t = next((x for x in project.get("tables", []) if x["id"] == a["table"]), None)
        if t is None:
            continue
        name = name_of(a["name"])
        try:
            res = analysis.run(a["kind"], t, a.get("sets", []), a.get("options", {}))
        except Exception as e:
            out.append((name, [[a["name"]], [f"Could not calculate: {e}"]], True))
            continue
        rows = [[(res["title"], "title")], [res["answer"]], []]
        for tab in res["tabs"]:
            for tb in tab["tables"]:
                rows.append([(f"{tab['name']} · {tb['title']}" if tb["title"] else tab["name"], "bold")])
                rows.append([(c, "head") for c in tb["columns"]])
                rows += [[num(v) for v in row] for row in tb["rows"]]
                rows.append([])
        rows += [[("Methods", "bold")], [res["methods"]]]
        out.append((name, rows, True))
    return out


def write_excel(path, project, analyses):
    """File › Export › Workbook (G6): a sheet per data table and one per result, as .xlsx or .ods by the extension."""
    if path.lower().endswith(".ods"):
        from . import odf
        return odf.write_ods(path, sheets(project, analyses))
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    wb.remove(wb.active)
    look = {"head": {"font": Font(bold=True), "fill": PatternFill("solid", fgColor="E3F5F2")}, "bold": {"font": Font(bold=True)},
            "title": {"font": Font(bold=True, size=13)}, "excl": {"font": Font(italic=True, color="1F6FD0")}}
    for name, rows, wide in sheets(project, analyses):
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append([c[0] if isinstance(c, tuple) else c for c in row])
            for j, c in enumerate(row):
                if isinstance(c, tuple):
                    for k, v in look[c[1]].items():
                        setattr(ws.cell(ws.max_row, j + 1), k, v)
        if wide:
            ws.column_dimensions["A"].width = 34
        else:
            ws.freeze_panes = "A2"
    for ws in wb.worksheets:  # text such as "=HYPERLINK(…)" in a title must stay text, never become a live formula
        for row in ws.iter_rows():
            for c in row:
                if c.data_type == "f":
                    c.data_type = "s"
    with atomic(path) as tmp:
        wb.save(tmp)
    return {"path": path, "sheets": len(wb.sheetnames)}


def num(v):
    """Result cells are text; plain numbers go to Excel as numbers."""
    try:
        return float(v) if re.fullmatch(r"-?\d+(\.\d+)?(e[-+]?\d+)?", v.strip()) else v
    except (AttributeError, ValueError):
        return v

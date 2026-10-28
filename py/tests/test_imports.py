"""Importing (A13, A14): a small .pzfx project written as the format stores it (namespaced XML, an XY
table with replicates and an excluded value, a one-way table with row titles), a 96-well plate block
inside a sheet, and several run files at once. Run: python -m unittest discover -s py/tests"""
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from datum import imports, io  # noqa: E402

PZFX = """<?xml version="1.0" encoding="UTF-8"?>
<GraphPadPrismFile xmlns="http://graphpad.com/prism/Prism.htm" PrismXMLVersion="5.00">
<TableSequence><Ref ID="Table0"/><Ref ID="Table1"/></TableSequence>
<Table ID="Table0" XFormat="numbers" YFormat="replicates" Replicates="2" TableType="XY">
<Title>Growth</Title>
<XColumn Width="81" Subcolumns="1"><Title>Time</Title><Subcolumn><d>0</d><d>1</d><d>2</d></Subcolumn></XColumn>
<YColumn Width="162" Subcolumns="2"><Title>Control</Title>
<Subcolumn><d>1.0</d><d>2.0</d><d>4.1</d></Subcolumn><Subcolumn><d>1.1</d><d Excluded="1">9.9</d><d/></Subcolumn></YColumn>
</Table>
<Table ID="Table1" XFormat="none" YFormat="replicates" TableType="OneWay">
<Title>Strength</Title>
<RowTitlesColumn Width="40"><Subcolumn><d>s1</d><d>s2</d></Subcolumn></RowTitlesColumn>
<YColumn Subcolumns="1"><Title>Annealed</Title><Subcolumn><d>210</d><d>215</d></Subcolumn></YColumn>
<YColumn Subcolumns="1"><Title>Quenched</Title><Subcolumn><d>320</d><d>n/a</d></Subcolumn></YColumn>
</Table>
</GraphPadPrismFile>"""


def write(dirname, name, text):
    p = os.path.join(dirname, name)
    open(p, "w", encoding="utf-8").write(text)
    return p


class Imports(unittest.TestCase):
    def test_pzfx_tables(self):
        with tempfile.TemporaryDirectory() as d:
            r = io.read_file(write(d, "p.pzfx", PZFX))
        a, b = r["tables"]
        self.assertEqual((a["name"], a["kind"], a["xTitle"], a["sets"], a["rows"]), ("Growth", "xy", "Time", [{"title": "Control", "entry": "rep", "reps": 2}], 3))
        self.assertEqual(a["cols"], [[0, 1, 2], [1.0, 2.0, 4.1], [1.1, 9.9, None]])
        self.assertEqual(a["excluded"], [[1, 2]])
        self.assertEqual((b["kind"], [s["title"] for s in b["sets"]], b["rowTitles"], b["cols"][2]), ("column", ["Annealed", "Quenched"], ["s1", "s2"], [320.0, None]))

    def test_plate_block_in_a_sheet(self):
        rows = [["Plate 1", "", ""], ["", *[str(i) for i in range(1, 13)]]] + [[chr(65 + r), *[str(r * 12 + c) for c in range(12)]] for r in range(8)] + [["end"]]
        p = imports.plates([{"name": "Read", "rows": rows}])
        self.assertEqual(len(p), 1)
        t = p[0]
        self.assertEqual((t["kind"], t["rows"], len(t["sets"]), t["rowTitles"][7], t["cols"][1][1], t["cols"][12][7]), ("grouped", 8, 12, "H", 12.0, 95.0))
        with tempfile.TemporaryDirectory() as d:
            self.assertEqual(len(io.read_file(write(d, "plate.csv", "\n".join(",".join(r) for r in rows)))["plates"]), 1)

    def test_several_runs(self):
        with tempfile.TemporaryDirectory() as d:
            a = write(d, "run1.csv", "Strain,Stress\n0,0\n0.1,50\n0.2,90\n")
            b = write(d, "run2.csv", "Strain,Stress\n0,0\n0.1,55\n0.2,95\n")
            c = write(d, "other.csv", "t,y\n0,1\n5,2\n")
            one = imports.read_many([a, b])
            self.assertEqual((len(one), one[0]["xTitle"], [s["title"] for s in one[0]["sets"]], one[0]["cols"][2]), (1, "Strain", ["run1", "run2"], [0.0, 55.0, 95.0]))
            two = imports.read_many([a, c])
            self.assertEqual([t["name"] for t in two], ["run1", "other"])


class Reports(unittest.TestCase):
    def test_slides_and_word_report(self):
        from test_groups import col
        from datum import report
        from docx import Document
        from pptx import Presentation
        t = {**col([4.2, 4.6, 4.4, 4.9], [5.6, 5.9, 5.4, 6.1]), "id": 1, "name": "Yield"}
        p = {"name": "P", "info": {"title": "Heat treatment", "experiment": "Two tempers."}, "tables": [t],
             "graphs": [{"id": 2, "table": 1, "name": "Yield graph", "settings": {"column": "scatter", "points": "each", "overlays": [3]}}],
             "analyses": [{"id": 3, "table": 1, "name": "t-test", "kind": "ttest", "sets": [], "options": {}}]}
        with tempfile.TemporaryDirectory() as d:
            s = report.slides(p, os.path.join(d, "s.pptx"))
            self.assertEqual((s["slides"], len(Presentation(os.path.join(d, "s.pptx")).slides)), (3, 3))
            w = report.word(p, os.path.join(d, "r.docx"))
            doc = Document(os.path.join(d, "r.docx"))
            text = "\n".join(x.text for x in doc.paragraphs)
            self.assertEqual((w["figures"], w["results"], len(doc.inline_shapes)), (1, 1, 1))
            for words in ("Heat treatment", "Figure 1. Yield graph", "t-test", "Methods"):
                self.assertIn(words, text)
            self.assertGreater(len(doc.tables[0].rows), 1)

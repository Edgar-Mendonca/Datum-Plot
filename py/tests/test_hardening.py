"""The hardening standard for the Python engine (FAMILY §11): damaged files for every reader give a clear
error; the worker answers every bad request with an error and keeps going; exports are atomic and
spreadsheets never carry live formulas; random tables never make an analysis or a graph crash.
Run: python -m unittest discover -s py/tests"""
import json
import os
import random
import subprocess
import sys
import tempfile
import unittest
import warnings

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from datum import analysis as A, files, graph, io as dio  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


class DamagedFiles(unittest.TestCase):
    def test_every_reader_opens_or_says_why(self):
        rnd = random.Random(5)
        with tempfile.TemporaryDirectory() as d:
            cases = {"a.csv": b"x,y\n1,2\n\x00\xff\xfe,\x81\n" + b"9," * 5000, "b.txt": bytes(rnd.randrange(256) for _ in range(4000)), "c.tsv": b"",
                     "d.xlsx": b"PK\x03\x04garbage", "e.xlsx": bytes(rnd.randrange(256) for _ in range(3000)), "f.pzfx": b"<GraphPadPrismFile><Table", "g.pzfx": b"\xff\xfe<<<",
                     "h.xlsm": b"", "i.ods": b"x"}
            for name, data in cases.items():
                p = os.path.join(d, name)
                open(p, "wb").write(data)
                try:
                    r = dio.read_file(p)
                    self.assertIsInstance(r["sheets"], list, name)
                except ValueError as e:
                    self.assertTrue(str(e), name)

    def test_the_worker_answers_bad_requests_and_keeps_going(self):
        lines = ["not json", "{}", '{"cmd": "nope"}', '{"cmd": "analyze", "kind": "linear"}', '{"cmd": "render", "table": 5, "graph": []}',
                 json.dumps({"cmd": "console", "project": {"tables": []}, "code": "x" * 300000}), '{"cmd": "ping"}']
        out = subprocess.run([sys.executable, "-W", "ignore", os.path.join(ROOT, "worker.py")], input="\n".join(lines) + "\n", capture_output=True, text=True, timeout=300).stdout.splitlines()
        replies = [json.loads(line) for line in out]
        self.assertTrue(replies[0].get("ready"))
        self.assertEqual([r["ok"] for r in replies[1:]], [False, False, False, False, False, True, True])
        self.assertIn("numpy", replies[-1])


class Exports(unittest.TestCase):
    def test_spreadsheets_carry_no_live_formulas_and_saves_are_atomic(self):
        from openpyxl import load_workbook
        t = {"id": 1, "name": "=HYPERLINK(\"http://x\")", "kind": "xy", "xTitle": "+cmd", "rows": 2, "rowTitles": ["@SUM(A1)", ""], "excluded": [],
             "sets": [{"title": "=1+1", "entry": "rep", "reps": 1}], "cols": [[1, 2], [3, 4]]}
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "x.xlsx")
            dio.write_excel(p, {"name": "P", "tables": [t]}, [])
            wb = load_workbook(p)
            cells = [c for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None]
            self.assertTrue(cells and not [c.value for c in cells if c.data_type == "f"])
            self.assertIn("=1+1", [c.value for c in cells])
            self.assertEqual(sorted(os.listdir(d)), ["x.xlsx"], "no temporary file left")
            with self.assertRaises(ZeroDivisionError):
                with files.atomic(os.path.join(d, "y.png")) as tmp:
                    open(tmp, "wb").write(b"half")
                    1 / 0
            self.assertEqual(sorted(os.listdir(d)), ["x.xlsx"], "a failed export leaves nothing behind")
            self.assertEqual([files.safe_text(v) for v in ("=1", "+a", "-2", "@x", "ok", 3)], ["'=1", "'+a", "'-2", "'@x", "ok", 3])


class RandomTables(unittest.TestCase):
    def test_random_tables_never_crash_an_analysis_or_a_graph(self):
        rnd = random.Random(42)
        kinds = list(A.more_kinds())
        tkinds = ["xy", "column", "grouped", "contingency", "parts", "multi", "survival", "nested", "sheet", "matrix"]
        warnings.simplefilter("ignore")
        for _ in range(25):
            n, k = rnd.randrange(1, 12), rnd.randrange(1, 4)
            reps = rnd.choice([1, 1, 2, 3])
            val = lambda: rnd.choice([None, 0.0, 1.0, -1.0, 1e-12, 1e12, rnd.gauss(5, 3), rnd.randrange(2)])
            t = {"name": "R", "kind": rnd.choice(tkinds), "rows": n, "xTitle": "X", "yTitle": "Y", "rowTitles": [""] * n, "excluded": [[0, 1]],
                 "sets": [{"title": f"S{i}", "entry": "rep", "reps": reps, "role": rnd.choice(["", "x", "yerr"])} for i in range(k)],
                 "cols": [[val() for _ in range(n)] for _ in range(1 + k * reps)]}
            for kind in kinds:
                try:
                    json.dumps(A.run(kind, t, [], {}))
                except ValueError:
                    pass
            graph.render(t, {"column": rnd.choice(["scatter", "bars", "box", "contour", "surface", "pie"]), "y": {"breakFrom": 1, "breakTo": 2}})


class OpenDocument(unittest.TestCase):
    def test_workbook_document_and_slides_are_valid_packages(self):
        import zipfile
        import xml.dom.minidom
        from datum import report
        t = {"id": 1, "name": "=cmd|' /C calc'!A0", "kind": "xy", "xTitle": "Time", "rows": 3, "rowTitles": [], "excluded": [[1, 1]],
             "sets": [{"title": "=1+1", "entry": "rep", "reps": 1}], "cols": [[1, 2, 3], [2.5, 4, None]]}
        a = {"id": 2, "table": 1, "name": "Fit <&>", "kind": "linear", "sets": [], "options": {}}
        p = {"name": "P", "info": {"title": "Test & check"}, "tables": [t], "analyses": [a], "graphs": [{"id": 3, "table": 1, "name": "G", "settings": {}}], "layouts": []}
        with tempfile.TemporaryDirectory() as d:
            dio.write_excel(os.path.join(d, "b.ods"), p, [a])
            report.word(p, os.path.join(d, "d.odt"))
            report.slides(p, os.path.join(d, "s.odp"))
            for f, mime in (("b.ods", "spreadsheet"), ("d.odt", "text"), ("s.odp", "presentation")):
                with zipfile.ZipFile(os.path.join(d, f)) as z:
                    first = z.infolist()[0]
                    self.assertEqual((first.filename, first.compress_type), ("mimetype", zipfile.ZIP_STORED), f)
                    self.assertEqual(z.read("mimetype").decode(), f"application/vnd.oasis.opendocument.{mime}")
                    for part in ("content.xml", "styles.xml", "META-INF/manifest.xml"):
                        xml.dom.minidom.parseString(z.read(part))  # well-formed
                    manifest = z.read("META-INF/manifest.xml").decode()
                    for n in z.namelist():
                        if n.startswith("Pictures/"):
                            self.assertIn(n, manifest)
                    content = z.read("content.xml").decode()
                    self.assertNotIn("table:formula", content, "text that looks like a formula stays text")
            self.assertIn("=1+1", zipfile.ZipFile(os.path.join(d, "b.ods")).read("content.xml").decode())
            self.assertEqual(sorted(os.listdir(d)), ["b.ods", "d.odt", "s.odp"], "no temporary files left")
